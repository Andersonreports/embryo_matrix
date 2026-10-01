"""Dry run: shows what the day/month date repair WOULD change in the dashboard's copy of the sheet rows. Writes nothing to the
database and never touches the Google Sheet. Usage: .venv/bin/python tools/date_fix_report.py [out.xlsx]"""
import collections, json, random, sqlite3, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl
from openpyxl.styles import Font, PatternFill
from app.date_fix import DATE_FIELDS, normalize_row

out = sys.argv[1] if len(sys.argv) > 1 else "date_fix_report.xlsx"
db = sqlite3.connect(Path(__file__).resolve().parent.parent / "embryomatrix.db")
rows = [r for r in json.loads(db.execute("select value from kv_store where key='embryomatrix-imported-cases'").fetchone()[0]) if not r.get("_stale")]
changed, unclear = [], []
per_field = collections.defaultdict(collections.Counter)
per_tab = collections.defaultdict(collections.Counter)
rows_changed = rows_unclear = 0
for r in rows:
    base = {**r, **(r.get("_rawDates") or {})}
    fixes, unc = normalize_row(base)
    for f in DATE_FIELDS:
        if base.get(f):
            per_field[f]["has a date"] += 1
    for f, new in fixes.items():
        per_field[f]["changed"] += 1
        per_tab[r.get("_importSource") or "?"]["dates changed"] += 1
        changed.append([r.get("_importSource"), r.get("patient name"), r.get("sample id"), r.get("embryo name"), f, base.get(f), new, base.get("date sample received") if f != "date sample received" else ""])
    for f in unc:
        per_field[f]["unclear"] += 1
        per_tab[r.get("_importSource") or "?"]["dates unclear"] += 1
        unclear.append([r.get("_importSource"), r.get("patient name"), r.get("sample id"), r.get("embryo name"), f, base.get(f)])
    rows_changed += bool(fixes); rows_unclear += bool(unc)
    per_tab[r.get("_importSource") or "?"]["rows"] += 1
wb = openpyxl.Workbook()
def sheet(name, head, data, widths):
    ws = wb.create_sheet(name); ws.append(head)
    for row in data: ws.append(row)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0A7180")
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
s = wb.active; s.title = "Summary"
s.append(["Rows checked", len(rows)]); s.append(["Rows with at least one date changed", rows_changed]); s.append(["Rows with at least one unclear date", rows_unclear]); s.append([])
s.append(["Date column", "Has a date", "Would change", "Unclear (left as is)"])
for f in DATE_FIELDS: s.append([f, per_field[f]["has a date"], per_field[f]["changed"], per_field[f]["unclear"]])
s.append([]); s.append(["Sheet", "Rows", "Dates changed", "Dates unclear"])
for t in sorted(per_tab): s.append([t, per_tab[t]["rows"], per_tab[t]["dates changed"], per_tab[t]["dates unclear"]])
s.column_dimensions["A"].width = 38
for c in "BCD": s.column_dimensions[c].width = 22
random.seed(7)
recv = [c for c in changed if c[4] == "date sample received"]
by_tab = collections.defaultdict(list)
for c in recv: by_tab[c[0]].append(c)
check = []
for t, lst in by_tab.items(): check += random.sample(lst, min(4, len(lst)))
random.shuffle(check)
sheet("Check these (spot check)", ["Sheet", "Patient", "Sample ID", "Embryo", "Column", "Stored in sheet", "Would become", "Received (for context)"], check, [16, 28, 20, 18, 22, 18, 18, 18])
sheet("All changes", ["Sheet", "Patient", "Sample ID", "Embryo", "Column", "Stored in sheet", "Would become", "Received"], changed, [16, 28, 20, 18, 22, 18, 18, 18])
sheet("Unclear (not changed)", ["Sheet", "Patient", "Sample ID", "Embryo", "Column", "Stored in sheet"], unclear, [16, 28, 20, 18, 22, 22])
wb.save(out)
print("rows", len(rows), "| rows changed", rows_changed, "| rows with unclear", rows_unclear, "| dates changed", len(changed), "| unclear dates", len(unclear))
for f in DATE_FIELDS: print(f.ljust(22), dict(per_field[f]))
print("saved", out)
