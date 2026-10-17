# WGA and sequencing batches made on the Batch creation page. Each batch is kept in batches.json, written to the
# PGS & NIPGS updates (edits) sheet through the Apps Script (action batchBlockSave: one block per batch in the
# "WGA Batches" / "Sequencing Batches" tab, laid out like the live batch record sheets - Batch identification,
# Kit / reagent, QC check, Sample details) and downloaded as an .xlsx of the same block. The bench fills in the
# results (WGA concentrations, barcodes...) and uploads that file back: apply_results() reads it, updates the
# batch and re-writes its block. The live batch sheets are never written to.
import base64
import io
import json
from datetime import date
import re
import threading
from datetime import datetime
from pathlib import Path
from threading import Lock

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import edits_sheet

LOG_PATH = Path(__file__).parent / "data" / "batches.json"
_lock = Lock()

KINDS = {"wga": "WGA Batches", "seq": "Sequencing Batches"}
# Batch-level fields per kind: (key, label in the sheet)
META = {
    "wga": [("batchNo", "WGA Batch No"), ("date", "Processing Date"), ("instrument", "Instrument Name"), ("instrumentId", "Instrument ID")],
    "seq": [("batchNo", "Run ID"), ("date", "Run Date"), ("seqType", "Seq Type"), ("instrument", "Seq Platform"),
            ("instrumentId", "Instrument ID"), ("operator", "Operator")],
}
KIT = [("process", "Process"), ("name", "Kit / Reagent"), ("manufacturer", "Manufacturer"), ("lot", "Lot No."),
       ("expiry", "Expiry Date"), ("qty", "Qty Used"), ("preparedBy", "Prepared By"), ("checkedBy", "Checked By")]
# Sample table per kind: (key, column header). Keys the page doesn't fill stay blank for the bench to complete.
SAMPLE_COLS = {
    "wga": [("sno", "S. No."), ("biopsy", "Biopsy Date"), ("received", "Received Date"), ("sampleType", "Sample Type"),
            ("patient", "Patient Name"), ("embryo", "Embryo Tag"), ("index", "Barcode / Index"),
            ("dnaConc", "DNA Conc. (ng/µL)"), ("purified", "Purified WGA Conc. (ng/µL)")],
    # As in the live Sequencing Batch Record sheet's Sample Processing Details (run 101A)
    "seq": [("sno", "S. No."), ("received", "Received Date"), ("patient", "Patient Name"), ("embryo", "Embryo Tag"),
            ("dnaConc", "DNA Conc. (ng/µL)"), ("volume", "Volume"), ("pool", "pool number"), ("index", "INDEX ID"),
            ("i7", "I7"), ("i5", "I5"), ("remarks", "Remarks")],
}


def _load() -> list:
    try:
        return json.loads(LOG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save(items: list) -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    LOG_PATH.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")


def _clean(v, n=200) -> str:
    return str(v if v is not None else "").strip()[:n]


def listing() -> list:
    return _load()


def create(kind: str, meta: dict, kit: dict, samples: list, user: str) -> dict:
    if kind not in KINDS:
        raise ValueError("Unknown batch kind")
    if not samples:
        raise ValueError("Choose at least one sample")
    meta = {k: _clean(meta.get(k)) for k, _ in META[kind]}
    if not meta["batchNo"]:
        raise ValueError(("WGA batch no" if kind == "wga" else "Run ID") + " is required")
    keys = [k for k, _ in SAMPLE_COLS[kind]] + ["sampleId", "test", "key", "editSid", "editTag"]
    rows = [{k: _clean(s.get(k)) for k in keys} for s in samples[:400]]
    for i, r in enumerate(rows):
        r["sno"] = str(i + 1)
    now = datetime.now()
    with _lock:
        items = _load()
        if any(b["kind"] == kind and b["meta"]["batchNo"].lower() == meta["batchNo"].lower() for b in items):
            raise ValueError(f"{'WGA batch' if kind == 'wga' else 'Run'} {meta['batchNo']} was already made here")
        entry = {"id": now.strftime("%Y%m%d%H%M%S%f"), "kind": kind, "meta": meta,
                 "kit": {k: _clean(kit.get(k)) for k, _ in KIT} if kit else {}, "samples": rows,
                 "createdBy": user, "createdAt": now.strftime("%Y-%m-%d %H:%M"), "sent": False, "error": ""}
        items.append(entry)
        _save(items)
    threading.Thread(target=retry_pending, daemon=True).start()
    return entry


def delete(batch_id: str) -> bool:
    with _lock:
        items = _load()
        hit = next((b for b in items if b["id"] == batch_id), None)
        if not hit:
            return False
        _save([b for b in items if b["id"] != batch_id])
    if hit.get("sent") and edits_sheet.configured():
        threading.Thread(target=_post_delete, args=(hit,), daemon=True).start()
    return True


def _post_delete(b: dict) -> None:
    try:
        edits_sheet._post({"action": "batchDelete", "tab": KINDS[b["kind"]], "id": b["id"]})
    except Exception:
        pass


def _body(b: dict) -> dict:
    """The batch's block for the edits sheet: the same grid, merges and styling as the .xlsx download."""
    lay = _layout(b)
    return {"action": "batchBlockSave", "tab": KINDS[b["kind"]], "id": b["id"], **lay,
            "logoB64": _logo_b64() if lay["logo"] else ""}


def retry_pending() -> int:
    """Send every unsent batch, oldest first; returns how many are still waiting."""
    if not edits_sheet.configured():
        return len([b for b in _load() if not (b.get("sent") and b.get("layout") == "block")])
    with _lock:
        items = _load()
        for b in items:
            # Batches sent before the block layout (one flat row per sample) are written again as a block.
            if b.get("sent") and b.get("layout") == "block":
                continue
            try:
                edits_sheet._post(_body(b))
                b["sent"], b["error"], b["layout"] = True, "", "block"
            except Exception as ex:
                b["error"] = str(ex)
                break
        _save(items)
        return len([b for b in items if not (b.get("sent") and b.get("layout") == "block")])


# ---- The batch block: one layout for the .xlsx download and the edits sheet ----
_THIN = Side(style="thin", color="000000")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUMERIC = ("sno", "dnaConc", "purified", "volume", "pool", "index")


def _num(v: str):
    return float(v) if re.fullmatch(r"\d+(\.\d+)?", v or "") else v


_BLUE, _PEACH = "CFE2F3", "FCE5CD"
LOGO_PATH = Path(__file__).resolve().parents[2] / "frontend" / "assets" / "logos" / "anderson-logo.png"
# Where each part of the block sits, per kind - the WGA record (8 columns wide) and the Sequencing Batch Record
# (11 wide, with the logo + title rows on top). Column spans are (first, last), 1-based.
_SPEC = {
    "wga": {"span": 8, "title": "", "ident_band": "BATCH IDENTIFICATION", "left": (1, 2), "label": (3, 4), "value": (5, 6),
            "right": (7, 8), "value_style": {}, "process": "WGA", "kit_at": [(1, 1), (2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7), (8, 8)],
            "sample_band": "SAMPLE DETAILS", "head_bg": "", "widths": [8, 14, 14, 20, 26, 12, 16, 16, 16]},
    "seq": {"span": 11, "title": "Sequencing Batch Record-Pre-Implantation Genetics", "ident_band": "", "left": (1, 3),
            "label": (4, 5), "value": (6, 9), "right": (10, 11), "value_style": {"bold": True, "size": 12, "center": True},
            "process": "SEQUENCING", "kit_at": [(1, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7), (8, 8), (9, 9)],
            "sample_band": "SAMPLE PROCESSING DETAILS", "head_bg": _PEACH, "widths": [8, 20, 26, 20, 20, 20, 20, 20, 19, 17, 22]},
}


def _layout(b: dict) -> dict:
    """The block as a grid: rows of values, merged ranges, styled ranges (background / bold / size / centred),
    bordered ranges and where the logo goes (all 1-based). Rendered by xlsx() here and by batchBlockSave_ in the
    Apps Script, so the download and the edits sheet look the same."""
    kind, cols = b["kind"], SAMPLE_COLS[b["kind"]]
    sp = _SPEC[kind]
    width, span = len(cols), _SPEC[kind]["span"]
    grid, merges, styles, boxes = {}, [], [], []

    def put(r, c, v):
        grid[(r, c)] = "" if v is None else v

    def style(r1, c1, r2, c2, **kw):
        styles.append({"r1": r1, "c1": c1, "r2": r2, "c2": c2, **kw})

    def band(r, text):
        merges.append([r, 1, r, span])
        put(r, 1, text)
        style(r, 1, r, span, bg=_BLUE, bold=True, size=13, center=True)
        boxes.append([r, r, span])

    def cell(r, at, v, **kw):
        c1, c2 = at
        if c2 > c1:
            merges.append([r, c1, r, c2])
        put(r, c1, v)
        if kw:
            style(r, c1, r, c2, **kw)

    r, logo = 1, None
    if sp["title"]:
        merges += [[1, 1, 3, 2], [1, 3, 3, span]]
        put(1, 3, sp["title"])
        style(1, 3, 3, span, bg=_PEACH, bold=True, size=16, center=True)
        boxes.append([1, 3, span])
        logo, r = [1, 1, 3, 2], 4
    if sp["ident_band"]:
        band(r, sp["ident_band"])
        r += 1
    meta = META[kind]
    top, bottom = r, r + len(meta) - 1
    for i, (k, label) in enumerate(meta):
        cell(top + i, sp["label"], label, bold=True, size=12)
        cell(top + i, sp["value"], b["meta"].get(k, ""), **sp["value_style"])
    merges += [[top, sp["left"][0], bottom, sp["left"][1]], [top, sp["right"][0], bottom, sp["right"][1]]]
    boxes.append([top, bottom, span])
    r = bottom + 1
    band(r, "KIT / REAGENT DETAILS")
    for (k, label), at in zip(KIT, sp["kit_at"]):
        cell(r + 1, at, label, bold=True, size=12, center=True)
        cell(r + 2, at, (b.get("kit") or {}).get(k, "") or (sp["process"] if k == "process" else ""), center=True)
    boxes.append([r + 1, r + 2, span])
    r += 3
    band(r, "QC CHECK")
    boxes.append([r + 1, r + 2, span])
    r += 3
    band(r, sp["sample_band"])
    for i, (_, label) in enumerate(cols):
        put(r + 1, i + 1, label)
    style(r + 1, 1, r + 1, width, bold=True, size=12, center=True, **({"bg": sp["head_bg"]} if sp["head_bg"] else {}))
    prev = None
    for n, s in enumerate(b["samples"]):
        row = r + 2 + n
        first = s.get("patient") != prev  # the sheets write the patient and dates on a patient's first embryo only
        prev = s.get("patient")
        for i, (k, _) in enumerate(cols):
            v = s.get(k, "")
            if k in ("patient", "biopsy", "received") and not first:
                v = ""
            put(row, i + 1, v)
    last_row = r + 1 + len(b["samples"])
    if b["samples"]:
        style(r + 2, 1, last_row, width, center=True)
    boxes.append([r + 1, last_row, width])
    full = max(width, span)
    rows = [[grid.get((rr, cc), "") for cc in range(1, full + 1)] for rr in range(1, last_row + 1)]
    return {"rows": rows, "merges": merges, "styles": styles, "boxes": boxes, "widths": sp["widths"][:full],
            "logo": logo, "sampleTop": r + 2}


def _logo_b64() -> str:
    try:
        return base64.b64encode(LOGO_PATH.read_bytes()).decode("ascii")
    except OSError:
        return ""


def xlsx(b: dict) -> tuple[bytes, str]:
    kind, lay = b["kind"], _layout(b)
    numeric_cols = {i + 1 for i, (k, _) in enumerate(SAMPLE_COLS[kind]) if k in _NUMERIC}
    wb = Workbook()
    ws = wb.active
    ws.title = ("WGA " if kind == "wga" else "Run ") + re.sub(r"[\\/*?:\[\]]", "-", b["meta"]["batchNo"])[:25]
    for r, row in enumerate(lay["rows"], 1):
        for c, v in enumerate(row, 1):
            if v != "":
                ws.cell(r, c, _num(v) if r >= lay["sampleTop"] and c in numeric_cols else v)
    for r1, c1, r2, c2 in lay["merges"]:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for r1, r2, c2 in lay["boxes"]:
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=1, max_col=c2):
            for c in row:
                c.border = _BOX
    for st in lay["styles"]:
        for row in ws.iter_rows(min_row=st["r1"], max_row=st["r2"], min_col=st["c1"], max_col=st["c2"]):
            for c in row:
                if st.get("bg"):
                    c.fill = PatternFill("solid", fgColor=st["bg"])
                if st.get("bold") or st.get("size"):
                    c.font = Font(bold=bool(st.get("bold")), size=st.get("size") or 11)
                if st.get("center"):
                    c.alignment = _CENTER
    for i, w in enumerate(lay["widths"]):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    if lay["logo"] and LOGO_PATH.exists():
        try:
            from openpyxl.drawing.image import Image as XlImage
            img = XlImage(str(LOGO_PATH))
            img.height, img.width = 54, round(54 * img.width / img.height)
            ws.add_image(img, "A1")
        except Exception:
            pass  # no Pillow: the title still says what the sheet is
    for r in range(1, len(lay["rows"]) + 1):
        ws.row_dimensions[r].height = 19
    ws.page_setup.orientation = "landscape"  # prints one page wide, like the record sheets
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    name = f"{'WGA_Batch' if kind == 'wga' else 'Seq_Run'}_{re.sub(r'[^A-Za-z0-9]+', '_', b['meta']['batchNo'])}.xlsx"
    return buf.getvalue(), name


# ---- Results uploaded back: the downloaded block, filled in at the bench ----
# Sample columns the bench fills in, per kind; a blank cell in the file leaves the saved value alone.
RESULT_KEYS = {"wga": ("dnaConc", "purified", "index", "sampleType"),
               "seq": ("dnaConc", "volume", "pool", "index", "i7", "i5", "remarks")}


def _txt(v, dotted=False) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y" if dotted else "%d-%m-%Y")
    if isinstance(v, float):
        return f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v).strip()


def _key(v) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(v or "").upper())


def _sample_col(header: str, kind: str):
    """Which sample field a header in the uploaded file is - the sheets word some differently
    ("WGA Conc." for "DNA Conc.", "Index" for "Barcode / Index")."""
    h = re.sub(r"\s+", " ", header.lower()).strip()
    if not h:
        return None
    if "purified" in h:
        return "purified"
    if "conc" in h:
        return "dnaConc"
    if "barcode" in h or h in ("index", "index id"):
        return "index"
    for k, label in SAMPLE_COLS[kind]:
        if h == label.lower() or h.rstrip(".") == label.lower().rstrip("."):
            return k
    return None


def read_results(data: bytes) -> dict:
    """Read a filled-in batch block: {kind, batchNo, meta, kit, samples: [{field: value}]}."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise ValueError("That isn't an Excel (.xlsx) file - upload the batch file you downloaded and filled in")
    ws = next((w for w in wb.worksheets if any(str(c.value or "").strip().upper() == "SAMPLE DETAILS"
                                               for row in w.iter_rows(max_col=3) for c in row)), wb.worksheets[0])
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    labels = {r: [str(v or "").strip() for v in row] for r, row in enumerate(rows)}
    flat = [x.lower() for row in labels.values() for x in row]
    kind = "wga" if "wga batch no" in flat else "seq" if "run id" in flat else ""
    if not kind:
        raise ValueError("Couldn't find 'WGA Batch No' or 'Run ID' in the file - is this a batch file from this page?")

    def right_of(label):
        for r, row in labels.items():
            for c, x in enumerate(row):
                if x.lower() == label.lower():
                    for v in rows[r][c + 1:]:
                        if _txt(v):
                            return _txt(v)
                    return ""
        return ""

    meta = {k: right_of(label) for k, label in META[kind]}
    kit = {}
    for r, row in labels.items():
        low = [x.lower() for x in row]
        if "kit / reagent" in low and "lot no." in low and r + 1 < len(rows):
            for k, label in KIT:
                if label.lower() in low:
                    kit[k] = _txt(rows[r + 1][low.index(label.lower())], dotted=(k == "expiry"))
            break
    head_r = next((r for r, row in labels.items() if "embryo tag" in [x.lower() for x in row]), None)
    if head_r is None:
        raise ValueError("Couldn't find the Sample Details header row (with 'Embryo Tag') in the file")
    cmap = {c: _sample_col(x, kind) for c, x in enumerate(labels[head_r])}
    samples, patient = [], ""
    for row in rows[head_r + 1:]:
        vals = {k: _txt(row[c]) for c, k in cmap.items() if k and c < len(row)}
        if not any(vals.values()):
            if samples:
                break
            continue
        if not vals.get("embryo"):
            continue
        patient = vals.get("patient") or patient  # the sheets write the patient on a patient's first embryo only
        vals["patient"] = patient
        samples.append(vals)
    return {"kind": kind, "batchNo": meta.get("batchNo", ""), "meta": meta, "kit": kit, "samples": samples}


def apply_results(data: bytes, user: str) -> dict:
    """Lay an uploaded, filled-in batch file over the batch made here with the same number. Returns the batch,
    the samples whose WGA concentrations changed (for the dashboard / edits sheet), and rows not matched."""
    got = read_results(data)
    kind, number = got["kind"], got["batchNo"]
    noun = "WGA batch" if kind == "wga" else "Run"
    with _lock:
        items = _load()
        b = next((x for x in items if x["kind"] == kind and _key(x["meta"]["batchNo"]) == _key(number)), None)
        if not b:
            raise ValueError(f"{noun} {number or '(no number in the file)'} wasn't made on this page - only batches created here can take results")
        for k, v in got["meta"].items():
            if v and k != "batchNo":
                b["meta"][k] = v[:200]
        for k, v in got["kit"].items():
            if v:
                b.setdefault("kit", {})[k] = v[:200]
        by_both = {(_key(s.get("patient")), _key(s.get("embryo"))): s for s in b["samples"]}
        by_tag = {}
        for s in b["samples"]:
            by_tag.setdefault(_key(s.get("embryo")), []).append(s)
        changed, unmatched, seen = [], [], set()
        for row in got["samples"]:
            s = by_both.get((_key(row.get("patient")), _key(row.get("embryo"))))
            if not s and len(by_tag.get(_key(row.get("embryo")), [])) == 1:
                s = by_tag[_key(row.get("embryo"))][0]
            if not s or id(s) in seen:
                unmatched.append(f"{row.get('patient', '')} {row.get('embryo', '')}".strip())
                continue
            seen.add(id(s))
            before = {k: s.get(k, "") for k in ("dnaConc", "purified")}
            for k in RESULT_KEYS[kind]:
                if row.get(k):
                    s[k] = row[k][:200]
            if kind == "wga" and any(s.get(k, "") != before[k] for k in before):
                changed.append(s)
        b["sent"], b["error"] = False, ""
        b["resultsAt"], b["resultsBy"] = datetime.now().strftime("%Y-%m-%d %H:%M"), user
        _save(items)
    threading.Thread(target=retry_pending, daemon=True).start()
    return {"batch": b, "changed": changed, "matched": len(seen), "unmatched": unmatched}
